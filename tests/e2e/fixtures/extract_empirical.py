"""Rebuilds minimal attributed E2E extracts from pinned source workbooks."""

from __future__ import annotations

import argparse
import csv
import hashlib
from collections import Counter
from collections.abc import Iterable
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

RETAIL_SHA256 = (
    "f5385cbb54bbebf7196389109c6b0621faab0c304e3702548165e71c84aede8b"
)
TRIAL_SHA256 = (
    "af33885a0754f3023abd658c61f2d70d0d6f51b0f9d2c4ff9e866097a71ce8d4"
)


def _verify(path: Path, digest: str) -> None:
    """Rejects source changes before deriving the pinned test oracle."""
    hasher = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            hasher.update(block)
    if hasher.hexdigest() != digest:
        raise ValueError(f"Unexpected source digest: {path}")


def _retail_rows(source: Path) -> list[tuple[str, ...]]:
    """Selects stable real invoices, including one cancellation per customer."""
    with ZipFile(source) as archive:
        with archive.open("Online Retail.xlsx") as workbook:
            sheet = load_workbook(
                workbook, read_only=True, data_only=True
            ).active
            if sheet is None:
                raise ValueError("Retail workbook has no active sheet")
            iterator = sheet.values
            if next(iterator) != (
                "InvoiceNo",
                "StockCode",
                "Description",
                "Quantity",
                "InvoiceDate",
                "UnitPrice",
                "CustomerID",
                "Country",
            ):
                raise ValueError("Unexpected retail columns")
            by_customer: dict[
                int,
                tuple[
                    list[tuple[int, tuple[object, ...]]],
                    list[tuple[int, tuple[object, ...]]],
                ],
            ] = {}
            for source_row, values in enumerate(iterator, start=2):
                if source_row > 50_001:
                    break
                customer = values[6]
                invoice = str(values[0]) if values[0] is not None else ""
                if not isinstance(customer, (int, float)) or not invoice:
                    continue
                sales, returns = by_customer.setdefault(int(customer), ([], []))
                selected = returns if invoice.startswith("C") else sales
                if not invoice.startswith("C") or not selected:
                    selected.append((source_row, values))
            eligible = sorted(
                customer
                for customer, (sales, returns) in by_customer.items()
                if len(sales) >= 9
                and len(returns) == 1
                and len({str(values[0]) for _, values in sales}) >= 2
            )[:24]
            if len(eligible) != 24:
                raise ValueError("Expected 24 eligible retail customers")
            output: list[tuple[str, ...]] = []
            for customer in eligible:
                sales, returns = by_customer[customer]
                first_per_invoice: dict[
                    str, tuple[int, tuple[object, ...]]
                ] = {}
                for sale in sales:
                    first_per_invoice.setdefault(str(sale[1][0]), sale)
                    if len(first_per_invoice) == 2:
                        break
                selected_sales = list(first_per_invoice.values())
                selected_rows = {source_row for source_row, _ in selected_sales}
                selected_sales.extend(
                    sale for sale in sales if sale[0] not in selected_rows
                )
                for source_row, values in sorted(
                    (*selected_sales[:9], *returns)
                ):
                    (
                        invoice,
                        stock,
                        description,
                        quantity,
                        date,
                        price,
                        _,
                        country,
                    ) = values
                    if not isinstance(date, datetime):
                        raise ValueError(
                            f"Invalid invoice date at row {source_row}"
                        )
                    output.append(
                        (
                            str(source_row),
                            str(customer),
                            str(invoice),
                            str(stock),
                            str(description or ""),
                            str(quantity),
                            date.isoformat(),
                            str(price),
                            str(country),
                        )
                    )
            return output


def _trial_rows(source: Path) -> list[tuple[str, ...]]:
    """Retains assignment and primary outcome without participant identifiers."""
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["MainData"]
    output: list[tuple[str, ...]] = []
    for source_row, values in enumerate(sheet.values, start=1):
        if source_row == 1 or values[1] not in (1, 6):
            continue
        if values[4] != "INCLUDE":
            continue
        outcome = values[3]
        if not isinstance(outcome, (int, float)) or outcome < 0:
            continue
        output.append((str(len(output) + 1), str(values[1]), str(outcome)))
    if Counter(row[1] for row in output) != {"1": 101, "6": 104}:
        raise ValueError("Unexpected trial sample sizes")
    return output


def _write(
    path: Path, header: Iterable[str], rows: Iterable[tuple[str, ...]]
) -> None:
    """Emits deterministic CSV with a fixed line ending and field order."""
    with path.open("w", encoding="utf-8", newline="") as output:
        writer = csv.writer(output, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)


def main() -> None:
    """Checks both originals and regenerates the committed minimal extracts."""
    parser = argparse.ArgumentParser()
    parser.add_argument("retail_zip", type=Path)
    parser.add_argument("trial_xlsx", type=Path)
    parser.add_argument("output_dir", type=Path)
    args = parser.parse_args()
    _verify(args.retail_zip, RETAIL_SHA256)
    _verify(args.trial_xlsx, TRIAL_SHA256)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    _write(
        args.output_dir / "retail.csv",
        (
            "source_row",
            "customer_id",
            "invoice_no",
            "stock_code",
            "description",
            "quantity",
            "invoice_date",
            "unit_price",
            "country",
        ),
        _retail_rows(args.retail_zip),
    )
    _write(
        args.output_dir / "trial.csv",
        ("record_id", "condition", "units_selected"),
        _trial_rows(args.trial_xlsx),
    )


if __name__ == "__main__":
    main()
